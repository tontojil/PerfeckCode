"""Template Informe Maestro v1 — Formato neutro profesional para informes universales.
Nivel Diente Feliz: portada sobria, indices automaticos, tablas numeradas con Nota, cierres por capitulo.
Uso: importar y llamar funciones. NO copiar este codigo.
"""
from docx import Document
from docx.shared import Cm, Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import parse_xml
import os

# ── Constantes neutras profesionales ──
PRIMARIO = "0F766E"      # Verde azulado sobrio (encabezados H2, acentos)
TEXTO_OSCURO = "1F2937"  # Gris oscuro texto y H1
GRIS_MEDIO = "4C4C4C"    # Subtitulos y metadatos
GRIS_CLARO = "F8FAFC"    # Fondo alterno suave
BORDE_TABLA = "BFBFBF"   # Bordes tabla
HEADER_BG = "1F2937"     # Fondo encabezado tabla
HEADER_TEXT = "FFFFFF"
FILA_ALT = "EDEFF2"      # Fila alternada gris muy claro
FOOTER_TEXT = "666666"
FIGURA_GRIS = "595959"

MARGIN_CM = 2.5


def create_doc(titulo, autores, institucion="", asignatura="",
               docente="", fecha="", subtitulo=""):
    """Crea documento con portada neutra y seccion de cuerpo con footer numerado.

    Retorna doc. Portada sin numero, cuerpo con numero de pagina.
    """
    doc = Document()

    s1 = doc.sections[0]
    _setup_page(s1)
    _enable_title_page(s1)

    # Portada centrada verticalmente con espacios
    for _ in range(8):
        _empty(doc)
    _center(doc, titulo, Pt(18), bold=True, color=TEXTO_OSCURO)
    if subtitulo:
        _center(doc, subtitulo, Pt(13), color=GRIS_MEDIO)
    _empty(doc)
    for nombre in autores:
        _center(doc, nombre, Pt(12), color=GRIS_MEDIO)
    _empty(doc)
    if institucion:
        _center(doc, institucion, Pt(11), color=GRIS_MEDIO)
    if asignatura:
        _center(doc, asignatura, Pt(11), color=GRIS_MEDIO)
    if docente:
        _center(doc, docente, Pt(11), color=GRIS_MEDIO)
    _empty(doc)
    if fecha:
        _center(doc, fecha, Pt(11), color=FOOTER_TEXT)

    # Seccion cuerpo con footer
    s2 = doc.add_section()
    _setup_page(s2)
    _remove_title_page(s2)
    s2.footer.is_linked_to_previous = False
    _setup_footer(s2.footer, asignatura or institucion or "Informe")

    _setup_styles(doc)
    return doc


def heading(doc, text, level=1):
    """Agrega heading jerarquico. H1 con linea divisoria color primario."""
    h = doc.add_heading(text, level=level)
    if level == 1:
        pPr = h._p.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<w:bottom w:val="single" w:sz="8" w:space="4" w:color="{PRIMARIO}"/>'
            f'</w:pBdr>'
        )
        pPr.append(pBdr)
    return h


def body(doc, text):
    """Parrafo cuerpo justificado 11pt interlineado 1.15."""
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
    return p


def cierre(doc, text):
    """Parrafo de cierre de capitulo con estilo distinguido (cursiva + borde izquierdo)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:left w:val="single" w:sz="8" w:space="6" w:color="{PRIMARIO}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)
    run = p.add_run(f"Cierre del capitulo. {text}")
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.italic = True
    return p


def section(doc, title, paragraphs, level=1, cierre_text=None):
    """Heading + parrafos + cierre opcional."""
    heading(doc, title, level=level)
    for text in paragraphs:
        body(doc, text)
    if cierre_text:
        cierre(doc, cierre_text)


def bullet(doc, text):
    """Bullet simple 11pt."""
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.clear()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    return p


def bullets(doc, items):
    """Lista de bullets."""
    for item in items:
        bullet(doc, item)


def tabla_maestra(doc, numero, titulo, headers, rows, fuente="Elaboracion propia."):
    """Tabla numerada estilo Diente Feliz con titulo arriba y Nota abajo.

    Genera: 'Tabla N' en negrita + titulo, luego tabla, luego Nota.
    """
    p_num = doc.add_paragraph()
    p_num.paragraph_format.space_before = Pt(8)
    p_num.paragraph_format.space_after = Pt(2)
    r1 = p_num.add_run(f"Tabla {numero}. {titulo}")
    r1.bold = True
    r1.font.name = "Calibri"
    r1.font.size = Pt(11)

    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        _set_cell_borders(cell, BORDE_TABLA)
        _shade_cell(cell, HEADER_BG)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(htext)
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row_data in enumerate(rows):
        bg = "FFFFFF" if ri % 2 == 0 else FILA_ALT
        for ci, cell_text in enumerate(row_data):
            cell = t.rows[ri + 1].cells[ci]
            _set_cell_borders(cell, BORDE_TABLA)
            _shade_cell(cell, bg)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(cell_text))
            run.font.size = Pt(10)
            run.font.name = "Calibri"

    p_nota = doc.add_paragraph()
    p_nota.paragraph_format.space_before = Pt(2)
    p_nota.paragraph_format.space_after = Pt(8)
    rn = p_nota.add_run(f"Nota. {fuente}")
    rn.italic = True
    rn.font.size = Pt(9)
    rn.font.name = "Calibri"
    rn.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    return t


def figura_pie(doc, numero, descripcion):
    """Pie de figura 9pt cursiva centrada."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(f"Figura {numero}. {descripcion}")
    run.font.size = Pt(9)
    run.italic = True
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    run.font.name = "Calibri"
    return p


def toc_entry(doc, level, text, page):
    """Entrada estilo Word con puntos y numero."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(
        '<w:tabs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:tab w:val="right" w:leader="dot" w:pos="9350"/>'
        '</w:tabs>'
    ))
    pf = p.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.0
    if level == 2:
        pf.left_indent = Cm(0.5)
    elif level == 3:
        pf.left_indent = Cm(1.0)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11) if level == 1 else Pt(10)
    r.bold = (level == 1)
    p.add_run("\t")
    rp = p.add_run(str(page))
    rp.font.name = "Calibri"
    rp.font.size = Pt(11) if level == 1 else Pt(10)
    rp.bold = (level == 1)
    return p


def insert_toc(doc):
    """Inserta UNICAMENTE la Tabla automatica de Word. Sin lista manual."""
    heading(doc, "Tabla de Contenido", level=1)
    p = doc.add_paragraph()
    run = p.add_run()
    run._r.append(parse_xml(
        '<w:fldChar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fldCharType="begin"/>'
    ))
    run._r.append(parse_xml(
        '<w:instrText xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>'
    ))
    run._r.append(parse_xml(
        '<w:fldChar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fldCharType="separate"/>'
    ))
    run2 = p.add_run("Haga clic derecho > Actualizar campo o vaya a Referencias > Actualizar tabla.")
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    run2.font.name = "Calibri"
    run2._r.append(parse_xml(
        '<w:fldChar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fldCharType="end"/>'
    ))
    doc.add_page_break()
    return p


def page_break(doc):
    """Salto de pagina."""
    doc.add_page_break()


def cite_apa(authors, year, page=None):
    """Cita parentetica APA 7. Retorna string."""
    if isinstance(authors, str):
        authors = [authors]
    n = len(authors)
    if n == 1:
        names = authors[0]
    elif n == 2:
        names = f"{authors[0]} & {authors[1]}"
    else:
        names = f"{authors[0]} et al."
    if page:
        return f"({names}, {year}, p. {page})"
    return f"({names}, {year})"


def cite_apa_narrative(authors, year, page=None):
    """Cita narrativa APA 7. Retorna string."""
    if isinstance(authors, str):
        authors = [authors]
    n = len(authors)
    if n == 1:
        names = authors[0]
    elif n == 2:
        names = f"{authors[0]} y {authors[1]}"
    else:
        names = f"{authors[0]} et al."
    if page:
        return f"{names} ({year}, p. {page})"
    return f"{names} ({year})"


def referencia_libro(doc, author, year, title, publisher, edition=None):
    """Referencia APA libro con sangria francesa."""
    ed = f" ({edition} ed.)" if edition else ""
    return _referencia(doc, f"{author} ({year}). {title}{ed}. {publisher}.")


def referencia_articulo(doc, author, year, title, journal, volume, issue=None, pages=None, doi=None):
    """Referencia APA articulo."""
    text = f"{author} ({year}). {title}. {journal}, {volume}"
    if issue:
        text += f"({issue})"
    if pages:
        text += f", {pages}"
    text += "."
    if doi:
        text += f" https://doi.org/{doi}"
    return _referencia(doc, text)


def referencia_web(doc, author, year, title, site, url):
    """Referencia APA web."""
    return _referencia(doc, f"{author} ({year}). {title}. {site}. {url}")


def referencia_norma(doc, emisor, year, titulo, url):
    """Referencia norma o ley chilena."""
    return _referencia(doc, f"{emisor} ({year}). {titulo}. {url}")


def save(doc, filepath):
    """Guarda creando directorios."""
    os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
    doc.save(filepath)


# ── Internas ──

def _setup_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(MARGIN_CM)
    section.bottom_margin = Cm(MARGIN_CM)
    section.left_margin = Cm(MARGIN_CM)
    section.right_margin = Cm(MARGIN_CM)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)


def _enable_title_page(section):
    sectPr = section._sectPr
    if sectPr.find(qn("w:titlePg")) is None:
        sectPr.append(parse_xml(
            '<w:titlePg xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
        ))


def _remove_title_page(section):
    sectPr = section._sectPr
    for el in sectPr.findall(qn("w:titlePg")):
        sectPr.remove(el)


def _setup_footer(footer, texto):
    p = footer.paragraphs[0]
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:top w:val="single" w:sz="4" w:space="4" w:color="{BORDE_TABLA}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)
    tabs = parse_xml(
        '<w:tabs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:tab w:val="right" w:pos="9360"/>'
        '</w:tabs>'
    )
    pPr.append(tabs)
    r1 = p.add_run(f"{texto}.")
    r1.font.size = Pt(9)
    r1.font.name = "Calibri"
    r1.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    tab_run = p.add_run()
    tab_run._r.append(parse_xml(
        '<w:tab xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
    ))
    r2 = p.add_run()
    r2.font.size = Pt(9)
    r2.font.name = "Calibri"
    r2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    r2._r.append(parse_xml(
        '<w:fldChar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fldCharType="begin"/>'
    ))
    r2._r.append(parse_xml(
        '<w:instrText xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xml:space="preserve"> PAGE </w:instrText>'
    ))
    r2._r.append(parse_xml(
        '<w:fldChar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fldCharType="end"/>'
    ))


def _setup_styles(doc):
    sn = doc.styles["Normal"]
    sn.font.name = "Calibri"
    sn.font.size = Pt(11)
    sn.paragraph_format.space_after = Pt(0)
    sn.paragraph_format.line_spacing = 1.15

    s1 = doc.styles["Heading 1"]
    s1.font.name = "Calibri"
    s1.font.size = Pt(14)
    s1.font.bold = True
    s1.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    s1.paragraph_format.space_before = Pt(12)
    s1.paragraph_format.space_after = Pt(4)

    s2 = doc.styles["Heading 2"]
    s2.font.name = "Calibri"
    s2.font.size = Pt(12)
    s2.font.bold = True
    s2.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
    s2.paragraph_format.space_before = Pt(10)
    s2.paragraph_format.space_after = Pt(4)

    s3 = doc.styles["Heading 3"]
    s3.font.name = "Calibri"
    s3.font.size = Pt(11)
    s3.font.bold = True
    s3.font.color.rgb = RGBColor(0x4C, 0x4C, 0x4C)
    s3.paragraph_format.space_before = Pt(8)
    s3.paragraph_format.space_after = Pt(2)


def _center(doc, text, size, bold=False, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.size = size
    run.bold = bold
    run.font.name = "Calibri"
    if color:
        r, g, b = int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16)
        run.font.color.rgb = RGBColor(r, g, b)
    return p


def _empty(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    return p


def _shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(
        f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="{color_hex}"/>'
    ))


def _set_cell_borders(cell, color_hex, sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(
        f'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'</w:tcBorders>'
    ))


def _referencia(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(
        '<w:ind xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:left="720" w:hanging="720"/>'
    ))
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = "Calibri"
    return p
